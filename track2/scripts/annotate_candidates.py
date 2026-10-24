#!/usr/bin/env python3
"""
Annotate Top10_helper candidates (evidence >= 3) from several public databases.

Usage
    python3 annotate_candidates.py Top10_helper_with_DB.tsv
    python3 annotate_candidates.py Top10_helper_with_DB.tsv --dry-run
    python3 annotate_candidates.py file.xlsx --min-evidence 4 --skip spliceai,literature

Input : Top10_helper as .tsv / .csv / .xlsx (columns "Evidence count (0-6)", "Variant", "Gene(s)").
Output: annotation_v2.tsv        one row per selected variant and gene
        mva_gene_overlap.tsv     every gene in the WHOLE sheet that Open Targets links to MVA
Needs : Python 3.8+ and internet (no extra packages; .xlsx input needs openpyxl).

Sources (each can be switched off with --skip gnomad,vep,spliceai,opentargets,literature)
  gnomad       gnomAD v4 API: per-population AF, homozygote count, gene constraint (pLI, LOEUF, missense z)
  vep          Ensembl VEP REST: MANE/canonical consequence, HGVS, exon, ClinVar significance
  spliceai     Broad SpliceAI lookup: max delta score (>=0.2 = predicted splice effect)
  opentargets  variant record, MVA association with per-source scores, reverse lookup of MVA genes
  literature   Europe PMC: papers mentioning the rsID, and gene + MVA papers
A failing source never stops the run: the error is written to the 'notes' column.
SVs/CNVs and variants without full alleles get gene-level annotation only.
"""
import argparse, csv, json, re, sys, time, urllib.parse, urllib.request, urllib.error

OT = "https://api.platform.opentargets.org/api/v4/graphql"
GNOMAD = "https://gnomad.broadinstitute.org/api"
VEP = "https://rest.ensembl.org/vep/human/region"
EPMC = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
SPLICEAI = "https://spliceai-lookup-4.broadinstitute.org/spliceai/"
VAR_RE = re.compile(r"\(chr([0-9XYM]+):(\d+) ([ACGT]+)>([ACGT]+)\)")
MVA_TERMS = re.compile(r"mosaic variegated|aneuploid|premature chromatid|chromosom\w* instab|segmented chromosome|"
                       r"atelis|microcephal", re.I)


# ---------------------------------------------------------------- http helpers
def http_json(url, data=None, headers=None, timeout=60, retries=3):
    h = {"Accept": "application/json", "User-Agent": "candidate-annotator/1.0"}
    h.update(headers or {})
    body = json.dumps(data).encode() if data is not None else None
    if body:
        h["Content-Type"] = "application/json"
    last = ""
    for i in range(retries):
        try:
            req = urllib.request.Request(url, body, h)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            last = f"HTTP {e.code}: {e.read()[:200].decode(errors='replace')}"
            if e.code < 500 and e.code != 429:
                break
        except Exception as e:
            last = str(e)
        time.sleep(2 * (i + 1))
    return {"_error": last}


def gql(api, query, variables=None):
    out = http_json(api, {"query": query, "variables": variables or {}})
    if "_error" in out:
        return None, out["_error"]
    if out.get("errors") and not out.get("data"):
        return None, out["errors"][0].get("message", "graphql error")[:200]
    return out.get("data"), (out["errors"][0]["message"][:120] if out.get("errors") else "")


# ------------------------------------------------- schema-aware query building
_sc = {}


def tfields(api, tname):
    """{field: {'type': base type name, 'args': [...]}} from GraphQL introspection (cached)."""
    key = (api, tname)
    if key in _sc:
        return _sc[key]
    d, _ = gql(api, """query($t:String!){__type(name:$t){fields{name args{name} type{name kind
        ofType{name kind ofType{name kind ofType{name kind}}}}}}}""", {"t": tname})
    f = {}
    for fld in ((d or {}).get("__type") or {}).get("fields") or []:
        ty = fld["type"]
        while ty and not ty.get("name") and ty.get("ofType"):
            ty = ty["ofType"]
        f[fld["name"]] = {"type": (ty or {}).get("name"), "args": [a["name"] for a in fld.get("args") or []]}
    _sc[key] = f
    return f


def build_sel(api, tname, wish):
    """Selection string from a wish list of (field, None | sub-wish-list); unknown fields are dropped."""
    tf = tfields(api, tname)
    parts = []
    for name, sub in wish:
        if name not in tf:
            continue
        if sub is None:
            parts.append(name)
        else:
            s = build_sel(api, tf[name]["type"], sub)
            if s:
                parts.append(f"{name}{{{s}}}")
    return " ".join(parts)


# ------------------------------------------------------------------- gnomAD
POP_WISH = [("id", None), ("ac", None), ("an", None), ("ac_hom", None)]
SEQ_WISH = [("ac", None), ("an", None), ("ac_hom", None), ("populations", POP_WISH)]
GN_VAR_WISH = [("variant_id", None), ("exome", SEQ_WISH), ("genome", SEQ_WISH)]
GN_GENE_WISH = [("symbol", None), ("gnomad_constraint",
                [("pli", None), ("oe_lof", None), ("oe_lof_upper", None), ("mis_z", None), ("lof_z", None)])]


def af_summary(v):
    """Combine exome + genome counts per population. Returns dict with verdict."""
    pops, ac_t, an_t, hom = {}, 0, 0, 0
    for seq in ("exome", "genome"):
        d = v.get(seq) or {}
        ac_t += d.get("ac") or 0; an_t += d.get("an") or 0; hom += d.get("ac_hom") or 0
        for p in d.get("populations") or []:
            pid = p.get("id") or ""
            if "_" in pid or pid == "remaining":
                continue  # sex splits and the 'remaining' bucket
            a = pops.setdefault(pid, [0, 0])
            a[0] += p.get("ac") or 0; a[1] += p.get("an") or 0
    afs = {k: a[0] / a[1] for k, a in pops.items() if a[1] >= 1000}
    best = max(afs.items(), key=lambda kv: kv[1]) if afs else ("", 0.0)
    gl = ac_t / an_t if an_t else 0.0
    if best[1] >= 0.05:
        verdict = "COMMON (>=5% in a population)"
    elif best[1] >= 0.01:
        verdict = "common (1-5%)"
    elif best[1] >= 0.001:
        verdict = "low-frequency (0.1-1%)"
    else:
        verdict = "rare (<0.1%)"
    return {"global_af": gl, "max_pop": f"{best[0]}:{best[1]:.2e}" if best[0] else "", "hom": hom,
            "verdict": verdict}


def gnomad_variant(vid):
    out = {"status": "", "global_af": "", "max_pop": "", "hom": "", "verdict": "", "err": ""}
    if not re.fullmatch(r"[0-9XYM]+-\d+-[ACGT]+-[ACGT]+", vid):
        out["err"] = "bad variant id"; return out
    root = tfields(GNOMAD, "Query").get("variant", {}).get("type")
    sel = build_sel(GNOMAD, root, GN_VAR_WISH) if root else ""
    if not sel:
        out["err"] = "gnomAD schema not readable"; return out
    d, err = gql(GNOMAD, '{variant(variantId:"%s", dataset:gnomad_r4){%s}}' % (vid, sel))
    if d is None:
        out["err"] = "gnomAD: " + err; return out
    v = d.get("variant")
    if not v:
        out.update(status="absent from gnomAD v4", verdict="absent (treat as rare)"); return out
    s = af_summary(v)
    out.update(status="present", global_af=f"{s['global_af']:.2e}", max_pop=s["max_pop"],
               hom=s["hom"], verdict=s["verdict"])
    return out


_gn_gene = {}


def gnomad_constraint(sym):
    if sym in _gn_gene:
        return _gn_gene[sym]
    out = {"pli": "", "loeuf": "", "mis_z": "", "err": ""}
    if re.fullmatch(r"[A-Za-z0-9._-]+", sym or ""):
        root = tfields(GNOMAD, "Query").get("gene", {}).get("type")
        sel = build_sel(GNOMAD, root, GN_GENE_WISH) if root else ""
        if sel:
            d, err = gql(GNOMAD, '{gene(gene_symbol:"%s", reference_genome:GRCh38){%s}}' % (sym, sel))
            g = (d or {}).get("gene")
            c = (g or {}).get("gnomad_constraint") or {}
            if c:
                f = lambda k: "" if c.get(k) is None else f"{c[k]:.3g}"
                out.update(pli=f("pli"), loeuf=f("oe_lof_upper"), mis_z=f("mis_z"))
            elif d is None:
                out["err"] = "gnomAD constraint: " + err
    _gn_gene[sym] = out
    return out


# ---------------------------------------------------------------------- VEP
def vep(chrom, pos, ref, alt):
    out = {"gene": "", "cons": "", "impact": "", "hgvsc": "", "hgvsp": "", "mane": "", "exon": "",
           "clinsig": "", "err": ""}
    r = http_json(VEP, {"variants": [f"{chrom} {pos} . {ref} {alt} . . ."], "hgvs": 1, "canonical": 1,
                        "mane": 1, "numbers": 1})
    if isinstance(r, dict) and "_error" in r:
        out["err"] = "VEP: " + r["_error"]; return out
    if not r:
        out["err"] = "VEP: empty"; return out
    rec = r[0]
    tcs = rec.get("transcript_consequences") or []
    pick = ([t for t in tcs if t.get("mane_select")] or [t for t in tcs if t.get("canonical") == 1] or tcs or [None])[0]
    if pick:
        out.update(gene=pick.get("gene_symbol", ""), cons=",".join(pick.get("consequence_terms") or []),
                   impact=pick.get("impact", ""), hgvsc=pick.get("hgvsc", ""), hgvsp=pick.get("hgvsp", ""),
                   mane=pick.get("mane_select", ""), exon=pick.get("exon", ""))
    else:
        out["cons"] = rec.get("most_severe_consequence", "")
    cs = set()
    for c in rec.get("colocated_variants") or []:
        cs.update(c.get("clin_sig") or [])
    out["clinsig"] = ",".join(sorted(cs))
    return out


# ------------------------------------------------------------------ SpliceAI
def spliceai(chrom, pos, ref, alt, base=SPLICEAI):
    q = urllib.parse.urlencode({"hg": 38, "distance": 500, "mask": 0, "variant": f"chr{chrom}-{pos}-{ref}-{alt}"})
    r = http_json(base + "?" + q, timeout=90, retries=2)
    if "_error" in r:
        return "", "SpliceAI: " + r["_error"]
    best = None
    try:
        for s in r.get("scores") or []:
            if isinstance(s, str):
                p = s.split("|")
                vals = [float(x) for x in p[2:6] if x not in ("", ".")]
            else:
                vals = [float(s[k]) for k in ("DS_AG", "DS_AL", "DS_DG", "DS_DL") if s.get(k) not in (None, "", ".")]
            if vals:
                best = max(vals + ([best] if best is not None else []))
    except Exception as e:
        return "", f"SpliceAI parse error: {e}"
    if best is None:
        return "", "SpliceAI: no score returned (" + str(r.get("error", ""))[:80] + ")"
    return f"{best:.2f}", ""


# --------------------------------------------------------------- Open Targets
OT_VAR_WISH = [("id", None), ("variantId", None), ("rsIds", None), ("hgvsId", None),
               ("mostSevereConsequence", [("label", None)]),
               ("alleleFrequencies", [("populationName", None), ("alleleFrequency", None)]),
               ("variantEffect", [("method", None), ("assessment", None), ("score", None)])]
OT_ROW_WISH = [("score", None), ("disease", [("name", None), ("id", None)]),
               ("datasourceScores", [("id", None), ("componentId", None), ("score", None)])]
OT_CON_WISH = [("constraintType", None), ("score", None), ("oe", None), ("oeUpper", None)]


def ot_variant(vid):
    out = {"found": "no", "rsids": "", "cons": "", "pred": "", "err": ""}
    q = tfields(OT, "Query").get("variant")
    if not q:
        out["found"] = "error"; out["err"] = "OT schema not readable"; return out
    sel = build_sel(OT, q["type"], OT_VAR_WISH)
    arg = q["args"][0] if q["args"] else "variantId"
    d, err = gql(OT, "query($v:String!){variant(%s:$v){%s}}" % (arg, sel), {"v": vid})
    if d is None:
        out["found"] = "error"; out["err"] = "OT variant: " + err; return out
    v = d.get("variant")
    if not v:
        return out
    out["found"] = "yes"
    out["rsids"] = ",".join(v.get("rsIds") or [])
    out["cons"] = (v.get("mostSevereConsequence") or {}).get("label", "")
    out["pred"] = "; ".join(f'{e.get("method")}:{e.get("assessment") or ""}:{e.get("score")}'
                            for e in v.get("variantEffect") or [])
    return out


_ot_gene = {}


def ot_gene(sym):
    if sym in _ot_gene:
        return _ot_gene[sym]
    res = {"id": "", "n": "", "mva": "none", "sources": "", "constraint": "", "err": ""}
    s, err = gql(OT, """query($q:String!){search(queryString:$q, entityNames:["target"], page:{index:0,size:5}){hits{id name entity}}}""",
                 {"q": sym})
    if s is None:
        res["err"] = "OT search: " + err; _ot_gene[sym] = res; return res
    hits = [h for h in (s["search"]["hits"] or []) if h["entity"] == "target"]
    pick = ([h for h in hits if (h["name"] or "").upper() == sym.upper()] or hits or [None])[0]
    if not pick:
        res["err"] = "OT: target not found"; _ot_gene[sym] = res; return res
    res["id"] = pick["id"]
    tf = tfields(OT, "Target")
    ad = tf.get("associatedDiseases", {}).get("type")
    rows_t = tfields(OT, ad).get("rows", {}).get("type") if ad else None
    rsel = build_sel(OT, rows_t, OT_ROW_WISH) if rows_t else "score disease{name}"
    csel = build_sel(OT, tf["geneticConstraint"]["type"], OT_CON_WISH) if "geneticConstraint" in tf else ""
    qq = ("query($id:String!){target(ensemblId:$id){associatedDiseases(page:{index:0,size:300}){count rows{%s}} %s}}"
          % (rsel, f"geneticConstraint{{{csel}}}" if csel else ""))
    d, err = gql(OT, qq, {"id": pick["id"]})
    if d is None:
        res["err"] = "OT target: " + err; _ot_gene[sym] = res; return res
    t = d["target"]
    rows = sorted(t["associatedDiseases"]["rows"], key=lambda r: -r["score"])
    res["n"] = t["associatedDiseases"]["count"]
    mva = [r for r in rows if MVA_TERMS.search(r["disease"]["name"]) and r["score"] >= 0.1]
    if mva:
        res["mva"] = "; ".join(f'{r["disease"]["name"]} ({r["score"]:.2f})' for r in mva[:4])
        src = mva[0].get("datasourceScores") or []
        res["sources"] = "; ".join(f'{x.get("id") or x.get("componentId") or "?"}:{(x.get("score") or 0):.2f}' for x in sorted(src, key=lambda x: -(x.get("score") or 0)))
    res["top1"] = f'{rows[0]["disease"]["name"]} ({rows[0]["score"]:.2f})' if rows else ""
    cons = t.get("geneticConstraint") or []
    res["constraint"] = "; ".join(f'{c.get("constraintType")}:oe={c.get("oe")}' + (f',upper={c.get("oeUpper")}' if c.get("oeUpper") is not None else "")
                                  for c in cons if c.get("constraintType") in ("lof", "syn", "mis"))
    _ot_gene[sym] = res
    return res


def ot_mva_targets():
    """All targets Open Targets links to MVA: {symbol: best score}."""
    s, err = gql(OT, """query{search(queryString:"mosaic variegated aneuploidy", entityNames:["disease"], page:{index:0,size:6}){hits{id name entity}}}""")
    if s is None:
        return {}, "OT disease search: " + err
    ids = [h["id"] for h in s["search"]["hits"] if h["entity"] == "disease" and re.search("mosaic variegated", h["name"] or "", re.I)]
    dt = tfields(OT, "Disease").get("associatedTargets", {}).get("type")
    rt = tfields(OT, dt).get("rows", {}).get("type") if dt else None
    sel = build_sel(OT, rt, [("score", None), ("target", [("approvedSymbol", None)])]) if rt else ""
    out = {}
    for i in ids[:4]:
        d, err = gql(OT, "query($id:String!){disease(efoId:$id){associatedTargets(page:{index:0,size:300}){rows{%s}}}}" % sel, {"id": i})
        for r in (((d or {}).get("disease") or {}).get("associatedTargets") or {}).get("rows") or []:
            sym = r["target"]["approvedSymbol"]
            out[sym] = max(out.get(sym, 0), r["score"])
    return out, ""


# ----------------------------------------------------------------- literature
def epmc(query, n=2):
    r = http_json(EPMC + "?" + urllib.parse.urlencode({"query": query, "format": "json", "pageSize": n, "resultType": "lite"}))
    if "_error" in r:
        return "", [], r["_error"]
    res = r.get("resultList", {}).get("result") or []
    return r.get("hitCount", 0), [f'{x.get("title","")[:90]} ({x.get("pubYear","")})' for x in res], ""


# ----------------------------------------------------------------------- main
def read_table(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8") as fh:
            rows = list(csv.reader(fh, delimiter="\t" if path.lower().endswith((".tsv", ".txt")) else ","))
    hdr = rows[0]
    for need in ("Evidence count (0-6)", "Variant", "Gene(s)"):
        if need not in hdr:
            sys.exit(f"Column '{need}' not found. Columns: {hdr}")
    ix = {h: hdr.index(h) for h in ("Evidence count (0-6)", "Variant", "Gene(s)")}
    return [dict(row=n, ev=int(r[ix["Evidence count (0-6)"]]), variant=r[ix["Variant"]].strip(), genes=r[ix["Gene(s)"]].strip())
            for n, r in enumerate(rows[1:], start=2) if r and r[ix["Evidence count (0-6)"]].strip().isdigit()]


def syms(s):
    return [g.strip() for g in s.split(",") if g.strip() and not g.strip().startswith("ENSG")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("table")
    ap.add_argument("--min-evidence", type=int, default=3)
    ap.add_argument("--skip", default="", help="comma list: gnomad,vep,spliceai,opentargets,literature")
    ap.add_argument("--spliceai-url", default=SPLICEAI)
    ap.add_argument("--out", default="annotation_v2.tsv")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    skip = {s.strip() for s in a.skip.split(",") if s.strip()}

    allrows = read_table(a.table)
    sel = [r for r in allrows if r["ev"] >= a.min_evidence]
    for r in sel:
        m = VAR_RE.search(r["variant"]) if "<" not in r["variant"] else None
        r["chrom"], r["pos"], r["ref"], r["alt"] = (m.groups() if m else ("", "", "", ""))
        r["vid"] = f"{r['chrom']}-{r['pos']}-{r['ref']}-{r['alt']}" if m else ""
        r["syms"] = syms(r["genes"])
    print(f"{len(sel)} variants with evidence >= {a.min_evidence} "
          f"({sum(1 for r in sel if r['vid'])} with full alleles, {sum(1 for r in sel if not r['vid'])} gene-level only)")
    if a.dry_run:
        for r in sel:
            print(f"  row {r['row']:>3} ev={r['ev']} {'variant+gene' if r['vid'] else 'gene only':12} {r['variant'][:70]}")
        return

    cols = ["sheet_row", "evidence", "variant", "gene",
            "gnomAD_status", "gnomAD_global_AF", "gnomAD_max_pop_AF", "gnomAD_hom_count", "AF_verdict",
            "OT_variant_found", "rsIDs", "OT_consequence", "OT_predictions",
            "VEP_gene", "VEP_consequence", "VEP_impact", "VEP_HGVSc", "VEP_HGVSp", "VEP_MANE", "VEP_exon", "VEP_ClinVar_sig",
            "SpliceAI_max_DS", "gene_pLI", "gene_LOEUF", "gene_missense_z",
            "OT_gene_id", "OT_n_associations", "OT_MVA_associations", "OT_MVA_sources", "OT_top_association",
            "OT_constraint", "PMC_rsID_hits", "PMC_gene_MVA_hits", "PMC_titles", "notes"]
    with open(a.out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t"); w.writerow(cols)
        for r in sel:
            notes = []
            g, o, v, sp, lit_rs = {}, {}, {}, "", ("", [])
            if r["vid"]:
                if "gnomad" not in skip:
                    g = gnomad_variant(r["vid"]); notes += [g["err"]] if g["err"] else []
                if "opentargets" not in skip:
                    o = ot_variant(r["vid"]); notes += [o["err"]] if o["err"] else []
                if "vep" not in skip:
                    v = vep(r["chrom"], r["pos"], r["ref"], r["alt"]); notes += [v["err"]] if v["err"] else []
                if "spliceai" not in skip:
                    sp, e = spliceai(r["chrom"], r["pos"], r["ref"], r["alt"], a.spliceai_url)
                    notes += [e] if e else []
                    time.sleep(1)
                if "literature" not in skip and o.get("rsids"):
                    rs = o["rsids"].split(",")[0]
                    n, titles, e = epmc(f'"{rs}"'); lit_rs = (n, titles); notes += [e] if e else []
            else:
                notes.append("gene-level only (SV/CNV or alleles not given)")
            for gene in (r["syms"] or [""]):
                c = gnomad_constraint(gene) if ("gnomad" not in skip and gene) else {}
                og = ot_gene(gene) if ("opentargets" not in skip and gene) else {}
                n2, t2 = "", []
                if "literature" not in skip and gene:
                    n2, t2, e = epmc(f'"{gene}" AND "mosaic variegated aneuploidy"')
                    notes_g = [e] if e else []
                else:
                    notes_g = []
                titles = (lit_rs[1] if lit_rs[1] else []) + t2
                w.writerow([r["row"], r["ev"], r["variant"], gene,
                            g.get("status", ""), g.get("global_af", ""), g.get("max_pop", ""), g.get("hom", ""), g.get("verdict", ""),
                            o.get("found", "n/a" if not r["vid"] else ""), o.get("rsids", ""), o.get("cons", ""), o.get("pred", ""),
                            v.get("gene", ""), v.get("cons", ""), v.get("impact", ""), v.get("hgvsc", ""), v.get("hgvsp", ""),
                            v.get("mane", ""), v.get("exon", ""), v.get("clinsig", ""),
                            sp, c.get("pli", ""), c.get("loeuf", ""), c.get("mis_z", ""),
                            og.get("id", ""), og.get("n", ""), og.get("mva", ""), og.get("sources", ""), og.get("top1", ""),
                            og.get("constraint", ""), lit_rs[0], n2, " || ".join(titles[:3]),
                            " | ".join([x for x in notes + notes_g + [og.get("err", ""), c.get("err", "")] if x])])
            print(f"done row {r['row']}: {r['variant'][:60]}")
    print("wrote", a.out)

    # reverse lookup: which genes anywhere in the sheet does Open Targets link to MVA?
    if "opentargets" not in skip:
        targets, err = ot_mva_targets()
        if err:
            print("reverse lookup failed:", err)
        else:
            by_gene = {}
            for r in allrows:
                for gs in syms(r["genes"]):
                    d = by_gene.setdefault(gs, {"n": 0, "best": 0, "variants": []})
                    d["n"] += 1; d["best"] = max(d["best"], r["ev"]); d["variants"].append(r["variant"][:50])
            hits = sorted(((s, targets[s], d) for s, d in by_gene.items() if s in targets), key=lambda x: -x[1])
            with open("mva_gene_overlap.tsv", "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh, delimiter="\t")
                w.writerow(["gene", "OT_MVA_score", "rows_in_sheet", "best_evidence_count", "example_variants"])
                for s, sc, d in hits:
                    w.writerow([s, f"{sc:.2f}", d["n"], d["best"], " ; ".join(d["variants"][:3])])
            print(f"wrote mva_gene_overlap.tsv ({len(hits)} sheet genes are MVA-linked in Open Targets, "
                  f"out of {len(targets)} MVA targets)")


if __name__ == "__main__":
    main()
